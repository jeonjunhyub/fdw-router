# -*- coding: utf-8 -*-
"""간단한 Markdown -> docx 변환기 (맑은 고딕, 표/제목/굵게/인라인코드 지원)"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = '맑은 고딕'
MONO = 'Consolas'


def set_font(run, name=FONT, size=None, bold=None, mono=False):
    n = MONO if mono else name
    run.font.name = n
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), n)
    rf.set(qn('w:hAnsi'), n)
    rf.set(qn('w:eastAsia'), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if mono:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`)')


def add_rich(par, text, size=10, base_bold=False):
    text = text.replace('\\|', '|')
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            set_font(par.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            set_font(par.add_run(part[1:-1]), size=size - 0.5, bold=base_bold, mono=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size=size, bold=base_bold); r.italic = True
        else:
            set_font(par.add_run(part), size=size, bold=base_bold)


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    out, cur, i = [], '', 0
    while i < len(line):
        if line[i] == '\\' and i + 1 < len(line) and line[i + 1] == '|':
            cur += '\\|'; i += 2; continue
        if line[i] == '|':
            out.append(cur.strip()); cur = ''; i += 1; continue
        cur += line[i]; i += 1
    out.append(cur.strip())
    return out


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), color)
    tcPr.append(sh)


def convert(src, dst):
    lines = open(src, encoding='utf-8').read().split('\n')
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if not s:
            i += 1; continue

        if s.startswith('```'):
            i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(6)
            set_font(p.add_run('\n'.join(buf)), size=8.5, mono=True)
            continue

        if re.match(r'^-{3,}$', s):
            i += 1; continue

        # 이미지: ![캡션](상대경로)  — md 파일 위치 기준으로 찾는다
        m = re.match(r'^!\[(.*?)\]\((.+?)\)\s*$', s)
        if m:
            cap, rel = m.group(1), m.group(2)
            img = os.path.join(os.path.dirname(os.path.abspath(src)), rel.replace('/', os.sep))
            if os.path.isfile(img):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img, width=Cm(16.5))
                if cap:
                    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    c.paragraph_format.space_after = Pt(8)
                    r = c.add_run(cap); set_font(r, size=9); r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            else:
                p = doc.add_paragraph(); add_rich(p, f'[이미지 없음: {rel}]', size=9)
            i += 1; continue

        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            lv = len(m.group(1)); txt = m.group(2)
            h = doc.add_heading(level=min(lv, 4))
            h.paragraph_format.space_before = Pt(14 if lv <= 2 else 10)
            h.paragraph_format.space_after = Pt(5)
            for r in list(h.runs):
                r._element.getparent().remove(r._element)
            sz = {1: 17, 2: 13.5, 3: 11.5, 4: 10.5}.get(lv, 10.5)
            add_rich(h, txt, size=sz, base_bold=True)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            i += 1; continue

        if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|?$', lines[i + 1].strip()):
            hdr = split_row(s); i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            n = len(hdr)
            t = doc.add_table(rows=1, cols=n)
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in enumerate(hdr):
                cell = t.rows[0].cells[c]
                cell.text = ''
                add_rich(cell.paragraphs[0], txt, size=9, base_bold=True)
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                shade(cell, 'EDEDED')
            for r in rows:
                cells = t.add_row().cells
                for c in range(n):
                    cells[c].text = ''
                    add_rich(cells[c].paragraphs[0], r[c] if c < len(r) else '', size=9)
                    cells[c].paragraphs[0].paragraph_format.space_after = Pt(2)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if s.startswith('> '):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip()); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_rich(p, ' '.join(x for x in buf if x), size=9.5)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            continue

        m = re.match(r'^(\s*)[-*]\s+(.*)$', ln)
        if m:
            lvl = len(m.group(1)) // 2
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(0.6 + 0.5 * lvl)
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(2))
            i += 1; continue

        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', ln)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(3))
            i += 1; continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, s)
        i += 1

    doc.save(dst)
    print('OK ->', dst)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
